"""Reports routes — lists available, issued, overdue books, and fine collections.
Supports exporting reports to PDF (ReportLab) and Excel (OpenPyXL).
"""

from flask import Blueprint, render_template, request, flash, redirect, url_for, send_file, abort
from flask_login import login_required, current_user
from datetime import date
import io

from extensions import mysql
from models.issued_book import IssuedBook
from models.book import Book

reports_bp = Blueprint('reports', __name__, url_prefix='/reports')

@reports_bp.before_request
@login_required
def limit_to_admin():
    if current_user.role != 'admin':
        abort(403)

@reports_bp.route('/')
@login_required
def index():
    """Render the reports dashboard with quick stats."""
    try:
        cur = mysql.connection.cursor()
        
        # Available books count
        cur.execute("SELECT COUNT(*) AS cnt FROM books WHERE available > 0")
        avail_count = cur.fetchone()['cnt']
        
        # Issued books count (active issues)
        cur.execute("SELECT COUNT(*) AS cnt FROM issued_books WHERE status = 'issued'")
        issued_count = cur.fetchone()['cnt']
        
        # Overdue count
        cur.execute("SELECT COUNT(*) AS cnt FROM issued_books WHERE status = 'issued' AND due_date < %s", (date.today(),))
        overdue_count = cur.fetchone()['cnt']
        
        # Total fines collected
        cur.execute("SELECT COALESCE(SUM(fine), 0) AS total FROM issued_books WHERE fine > 0")
        total_fines = float(cur.fetchone()['total'])
        
        cur.close()
        
        return render_template(
            'reports/index.html',
            avail_count=avail_count,
            issued_count=issued_count,
            overdue_count=overdue_count,
            total_fines=total_fines
        )
    except Exception as e:
        flash(f"Error loading reports dashboard: {e}", "danger")
        return render_template(
            'reports/index.html',
            avail_count=0,
            issued_count=0,
            overdue_count=0,
            total_fines=0.00
        )

def _get_report_data(report_type):
    """Retrieve raw list of data for a specific report type."""
    cur = mysql.connection.cursor()
    if report_type == 'available':
        cur.execute(
            "SELECT b.id, b.title, b.author, b.isbn, c.name AS category, b.available, b.shelf_location "
            "FROM books b "
            "LEFT JOIN categories c ON b.category_id = c.id "
            "WHERE b.available > 0 "
            "ORDER BY b.title ASC"
        )
        headers = ['ID', 'Title', 'Author', 'ISBN', 'Category', 'Available', 'Shelf Location']
        keys = ['id', 'title', 'author', 'isbn', 'category', 'available', 'shelf_location']
        title = "Available Books Report"
    elif report_type == 'issued':
        cur.execute(
            "SELECT ib.id, b.title AS book_title, s.name AS student_name, s.enrollment_no, "
            "       ib.issue_date, ib.due_date, ib.status "
            "FROM issued_books ib "
            "JOIN books b ON ib.book_id = b.id "
            "JOIN students s ON ib.student_id = s.id "
            "WHERE ib.status = 'issued' "
            "ORDER BY ib.issue_date DESC"
        )
        headers = ['Issue ID', 'Book Title', 'Student Name', 'Enrollment No', 'Issue Date', 'Due Date', 'Status']
        keys = ['id', 'book_title', 'student_name', 'enrollment_no', 'issue_date', 'due_date', 'status']
        title = "Currently Issued Books Report"
    elif report_type == 'overdue':
        cur.execute(
            "SELECT ib.id, b.title AS book_title, s.name AS student_name, s.enrollment_no, "
            "       ib.issue_date, ib.due_date, DATEDIFF(%s, ib.due_date) AS days_overdue "
            "FROM issued_books ib "
            "JOIN books b ON ib.book_id = b.id "
            "JOIN students s ON ib.student_id = s.id "
            "WHERE ib.status = 'issued' AND ib.due_date < %s "
            "ORDER BY ib.due_date ASC",
            (date.today(), date.today())
        )
        headers = ['Issue ID', 'Book Title', 'Student Name', 'Enrollment No', 'Issue Date', 'Due Date', 'Days Overdue']
        keys = ['id', 'book_title', 'student_name', 'enrollment_no', 'issue_date', 'due_date', 'days_overdue']
        title = "Overdue Books Report"
    elif report_type == 'fines':
        cur.execute(
            "SELECT ib.id, b.title AS book_title, s.name AS student_name, s.enrollment_no, "
            "       ib.return_date, ib.fine "
            "FROM issued_books ib "
            "JOIN books b ON ib.book_id = b.id "
            "JOIN students s ON ib.student_id = s.id "
            "WHERE ib.fine > 0 "
            "ORDER BY ib.return_date DESC"
        )
        headers = ['Issue ID', 'Book Title', 'Student Name', 'Enrollment No', 'Return Date', 'Fine (Rs.)']
        keys = ['id', 'book_title', 'student_name', 'enrollment_no', 'return_date', 'fine']
        title = "Fines Collection Report"
    else:
        headers = []
        keys = []
        title = "Report"
    
    rows = cur.fetchall()
    cur.close()
    return title, headers, keys, rows

@reports_bp.route('/view/<report_type>')
@login_required
def view_report(report_type):
    """View details of a specific report on a dedicated page."""
    try:
        title, headers, keys, rows = _get_report_data(report_type)
        return render_template(
            'reports/view.html',
            title=title,
            headers=headers,
            keys=keys,
            rows=rows,
            report_type=report_type
        )
    except Exception as e:
        flash(f"Error loading report: {e}", "danger")
        return redirect(url_for('reports.index'))

@reports_bp.route('/export/pdf/<report_type>')
@login_required
def export_pdf(report_type):
    """Export the specified report to a professional PDF file."""
    try:
        from reportlab.lib.pagesizes import letter
        from reportlab.lib import colors
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        
        title, headers, keys, rows = _get_report_data(report_type)
        
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=letter, rightMargin=30, leftMargin=30, topMargin=30, bottomMargin=30)
        story = []
        
        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            name='TitleStyle',
            parent=styles['Heading1'],
            fontSize=18,
            textColor=colors.HexColor('#1a1c2e'),
            spaceAfter=15
        )
        
        # Title
        story.append(Paragraph(title, title_style))
        story.append(Paragraph(f"Generated on: {date.today().strftime('%d-%b-%Y')}", styles['Normal']))
        story.append(Spacer(1, 15))
        
        # Prepare table data
        table_data = [headers]
        for row in rows:
            table_row = []
            for k in keys:
                val = row.get(k)
                if val is None:
                    table_row.append("")
                elif isinstance(val, date):
                    table_row.append(val.strftime('%d-%b-%Y'))
                elif isinstance(val, float):
                    table_row.append(f"{val:.2f}")
                else:
                    table_row.append(str(val))
            table_data.append(table_row)
        
        # Auto-compute column widths based on letter size (total width ~550pt)
        col_width = 550.0 / len(headers)
        
        t = Table(table_data, colWidths=[col_width]*len(headers))
        t.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#6366f1')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 9),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 6),
            ('BACKGROUND', (0, 1), (-1, -1), colors.HexColor('#f8fafc')),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#cbd5e1')),
            ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
            ('FONTSIZE', (0, 1), (-1, -1), 8),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ]))
        
        story.append(t)
        doc.build(story)
        buffer.seek(0)
        
        filename = f"{report_type}_report_{date.today().strftime('%Y%m%d')}.pdf"
        return send_file(
            buffer,
            as_attachment=True,
            download_name=filename,
            mimetype='application/pdf'
        )
    except Exception as e:
        flash(f"Error exporting PDF: {e}", "danger")
        return redirect(url_for('reports.index'))

@reports_bp.route('/export/excel/<report_type>')
@login_required
def export_excel(report_type):
    """Export the specified report to an Excel workbook using OpenPyXL."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        
        title, headers, keys, rows = _get_report_data(report_type)
        
        wb = Workbook()
        ws = wb.active
        ws.title = report_type.capitalize()
        
        # Header style
        header_font = Font(name='Arial', size=11, bold=True, color='FFFFFF')
        header_fill = PatternFill(start_color='6366F1', end_color='6366F1', fill_type='solid')
        align_left = Alignment(horizontal='left', vertical='center')
        
        # Add title row
        ws.append([title])
        ws.cell(row=1, column=1).font = Font(name='Arial', size=14, bold=True)
        ws.append([f"Generated on: {date.today().strftime('%d-%b-%Y')}"])
        ws.append([]) # Empty row
        
        # Add headers
        ws.append(headers)
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=4, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = align_left
        
        # Add rows
        for row in rows:
            row_data = []
            for k in keys:
                val = row.get(k)
                if isinstance(val, date):
                    row_data.append(val.strftime('%d-%b-%Y'))
                elif isinstance(val, float):
                    row_data.append(round(val, 2))
                else:
                    row_data.append(val)
            ws.append(row_data)
        
        # Auto-adjust column widths
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = col[0].column_letter
            ws.column_dimensions[col_letter].width = max(max_len + 3, 10)
            
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        
        filename = f"{report_type}_report_{date.today().strftime('%Y%m%d')}.xlsx"
        return send_file(
            buffer,
            as_attachment=True,
            download_name=filename,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
    except Exception as e:
        flash(f"Error exporting Excel: {e}", "danger")
        return redirect(url_for('reports.index'))
